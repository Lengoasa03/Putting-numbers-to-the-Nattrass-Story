"""
ECOH621 Part 4 deliverable: the counterfactual essay as a .docx.

Output: report/part4_counterfactual.docx
"""

import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pathlib import Path

# Paths resolve from this file's own location, so the project folder can be moved
# or copied anywhere (including onto a marker's machine) and still run.
PROJECT = Path(__file__).resolve().parents[2]   # .../Development project
RES = PROJECT / "resources"
BASE = PROJECT / "output"
OUT = BASE / "report" / "part4_counterfactual.docx"

TITLE = ("Part 4 \u2014 The Counterfactual: What Would Have Had to Be True for "
         "the LABOUR/ILO Strategy to Have Worked?")

INTRO = ("The LABOUR/ILO strategy was not merely \u2018spend more\u2019. It was a "
         "conditional bet, and the conditions were parameters. Four had to hold "
         "jointly.")

SECTIONS = [
    ("1. A large accelerator, and a demand impulse big enough to trigger it.",
     "The strategy required private investment to respond strongly to demand rather "
     "than to fiscal signals. Nattrass herself notes Chirinko\u2019s (1993:1883) "
     "finding that lagged demand is the single most significant empirical determinant "
     "of investment, so this was the better-supported prior even in 1996. My Part 3 "
     "estimates are consistent with it: excluding the 2009 and 2020 recessions, a "
     "one-point rise in lagged GDP growth is associated with a 0.75\u20130.94 point "
     "rise in real private investment growth (p < 0.01), while the deficit coefficient "
     "loses significance entirely. On this parameter the subsequent evidence broadly "
     "supports LABOUR/ILO \u2014 though the result is fragile to GDP vintage, and no "
     "coefficient here is causal."),
    ("2. Employment elasticity high enough to convert growth into jobs.",
     "Redistribution required demand-led growth to be labour-absorbing. The "
     "2001\u20132006 evidence is more favourable than the standard \u2018jobless "
     "growth\u2019 story allows: formal non-agricultural employment grew 1.9\u20134.5% "
     "a year through that expansion, briefly beating GEAR\u2019s own 4.3% target in "
     "2006. Demand-led growth did create formal jobs when it arrived. But Nattrass\u2019s "
     "warning cuts against reading too much into this: rising average labour "
     "productivity \u2014 pursued deliberately under \u2018high productivity now\u2019 "
     "industrial and wage policy \u2014 mechanically lowers the employment elasticity of "
     "output. The strategy needed labour-demanding growth while its own allies were "
     "bargaining for the wage floors that discourage it. Actual real wage growth in "
     "1996\u201399 (1.7, 2.3, 8.6, 3.0%) massively exceeded GEAR\u2019s assumed path "
     "(\u22120.5, 1.0, 1.0, 1.0%), and formal employment fell every one of those years. "
     "That is the internal contradiction Nattrass identified, and it is unresolved."),
    ("3. No punitive capital-account response.",
     "This is where Nattrass criticises LABOUR/ILO most sharply, and where the "
     "counterfactual is weakest. The strategy assumed policy autonomy that an open "
     "capital account does not grant. A deficit-financed expansion only works if the "
     "risk premium and capital outflow response to fiscal divergence are small. South "
     "Africa\u2019s subsequent history says they are not: the 2001 currency collapse, "
     "the 2008 reversal, and the post-2020 debt trajectory all show how quickly "
     "financing conditions tighten. Nattrass\u2019s contrast with the East Asian "
     "economies is decisive here \u2014 Malaysia and Thailand could run large deficits "
     "because they had mobilised deep domestic savings pools and were therefore "
     "\u2018less prone to destabilising capital flight\u2019. South Africa\u2019s savings "
     "rate was, and remains, low. The same deficit is simply a different instrument in "
     "a low-savings, open-capital-account economy."),
    ("4. A demand expansion that does not destabilise.",
     "Nattrass is explicit that the demand\u2013investment relationship has a limit: if "
     "expansion undermines macroeconomic balance through inflation and depreciation, "
     "profitability is threatened and investment falls rather than rises. Latin "
     "American macroeconomic populism is her counter-example. So LABOUR/ILO needed to "
     "be right not only that demand drives investment, but that it could deliver demand "
     "without triggering the instability that destroys it \u2014 a narrow path."),
]

VERDICT = ("The counterfactual is not \u2018LABOUR/ILO would have worked\u2019. It is "
           "that LABOUR/ILO was right about the mechanism GEAR got wrong \u2014 "
           "investment follows demand, not fiscal signalling \u2014 while being wrong, "
           "or at least unserious, about the constraint GEAR took seriously. "
           "Nattrass\u2019s criticism of both camps survives my own evidence intact: "
           "GEAR bet on a confidence channel my regressions cannot find, and LABOUR/ILO "
           "bet on policy space that globalisation and a low savings rate did not "
           "provide.")

REFERENCES = [
    "Chirinko, R. (1993) \u2018Business fixed investment spending: modeling strategies, "
    "empirical results and policy implications\u2019, Journal of Economic Literature 31. "
    "(Cited in Nattrass 2001.)",
    "Nattrass, N. (1996) \u2018Gambling on investment: competing economic strategies in "
    "South Africa\u2019, Transformation 31, 25\u201342.",
    "Nattrass, N. (2001) \u2018High productivity now: a critical review of South "
    "Africa\u2019s growth strategy\u2019, Transformation 45.",
    "South Africa (1996) Growth, Employment and Redistribution: a macroeconomic "
    "strategy. Pretoria: Department of Finance.",
]


def main():
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    h = doc.add_heading(TITLE, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    r = p.add_run(INTRO)
    r.font.size = Pt(11)

    word_total = len(INTRO.split())
    for head, text in SECTIONS:
        para = doc.add_paragraph()
        hr = para.add_run(head + " ")
        hr.bold = True
        hr.font.size = Pt(11)
        br = para.add_run(text)
        br.font.size = Pt(11)
        word_total += len(head.split()) + len(text.split())

    para = doc.add_paragraph()
    vr = para.add_run("Verdict. ")
    vr.bold = True
    vr.font.size = Pt(11)
    vt = para.add_run(VERDICT)
    vt.font.size = Pt(11)
    word_total += 1 + len(VERDICT.split())

    wc = doc.add_paragraph()
    wr = wc.add_run(f"({word_total} words, excluding title and references \u2014 "
                    f"within the 600-word limit.)")
    wr.italic = True
    wr.font.size = Pt(9)

    doc.add_heading("References", level=2)
    for ref in REFERENCES:
        rp = doc.add_paragraph()
        rr = rp.add_run(ref)
        rr.font.size = Pt(9.5)

    doc.save(OUT)
    print(f"Saved: {OUT}")
    print(f"Body word count: {word_total}")


if __name__ == "__main__":
    main()
