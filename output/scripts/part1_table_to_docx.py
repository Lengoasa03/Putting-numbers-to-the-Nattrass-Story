"""
ECOH621 Weeks 6-7 assignment -- Part 1 deliverable: GEAR-vs-reality table as a
.docx, styled after Nattrass's (2001) Table 1 "GEAR vs Reality".

Nattrass's original layout is two stacked panels -- "Predicted Results of GEAR"
and "Actual Performance" -- each with one row per indicator and one column per
year. This script reproduces that layout for the six indicators required by the
assignment brief, using the corrected/source-verified data produced by
part1_rebuild_table.py (run that script first).

Input : data/gear_target_vs_actual_1996_2000_CORRECTED.csv
Output: report/part1_gear_vs_reality_table.docx
"""

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from pathlib import Path

# Paths resolve from this file's own location, so the project folder can be moved
# or copied anywhere (including onto a marker's machine) and still run.
PROJECT = Path(__file__).resolve().parents[2]   # .../Development project
RES = PROJECT / "resources"
BASE = PROJECT / "output"
CORRECTED_CSV = BASE / "data" / "gear_target_vs_actual_1996_2000_CORRECTED.csv"
OUT_DOCX = BASE / "report" / "part1_gear_vs_reality_table.docx"

YEARS = [1996, 1997, 1998, 1999, 2000]

# Row order and display labels, matching Nattrass's Table 1 ordering as closely
# as the six required indicators allow.
ROWS = [
    ("real_gdp_growth_pct", "Real GDP growth (%)"),
    ("real_bank_rate_pct", "Real bank rate (%)"),
    ("real_private_investment_growth_pct", "Real private investment growth (%)"),
    ("nonagric_formal_employment_growth_pct", "Non-agric. formal employment growth (%)"),
    ("inflation_cpi_pct", "Inflation, CPI (%)"),
    ("deficit_pct_gdp", "Budget deficit (% of GDP)"),
]


def fmt(value) -> str:
    if pd.isna(value):
        return "\u2014"  # em dash, as Nattrass uses for missing 2000 actuals
    return f"{value:.1f}"


def shade_cell(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"


def build_panel(doc: Document, title: str, df: pd.DataFrame, value_col: str) -> None:
    heading = doc.add_paragraph()
    run = heading.add_run(title)
    run.bold = True
    run.italic = True
    run.font.size = Pt(11)

    table = doc.add_table(rows=len(ROWS) + 1, cols=len(YEARS) + 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    set_cell_text(table.rows[0].cells[0], "", bold=True)
    for j, year in enumerate(YEARS, start=1):
        set_cell_text(table.rows[0].cells[j], str(year), bold=True)
        shade_cell(table.rows[0].cells[j], "D9D9D9")
    shade_cell(table.rows[0].cells[0], "D9D9D9")

    # Data rows
    for i, (key, label) in enumerate(ROWS, start=1):
        set_cell_text(table.rows[i].cells[0], label, bold=False)
        table.rows[i].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_data = df[df["indicator"] == key].set_index("year")
        for j, year in enumerate(YEARS, start=1):
            val = row_data.loc[year, value_col] if year in row_data.index else pd.NA
            set_cell_text(table.rows[i].cells[j], fmt(val))

    # Column widths: wider label column, narrower year columns
    for row in table.rows:
        row.cells[0].width = Cm(6.0)
        for j in range(1, len(YEARS) + 1):
            row.cells[j].width = Cm(2.0)

    doc.add_paragraph()


def main() -> None:
    gear = pd.read_csv(CORRECTED_CSV)

    doc = Document()

    for section in doc.sections:
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    title = doc.add_heading(
        "Table 1: GEAR Targets versus Actual Performance, 1996\u20132000", level=1
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run(
        "Format after Nattrass, N. (2001) 'High Productivity Now', Table 1, "
        "'GEAR vs Reality'."
    )
    sub_run.italic = True
    sub_run.font.size = Pt(9)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    build_panel(doc, "Predicted Results of GEAR (Integrated Scenario)", gear, "gear_target")
    build_panel(doc, "Actual Performance", gear, "actual_to_1999")

    notes_heading = doc.add_paragraph()
    notes_heading.add_run("Notes and sources").bold = True

    notes = [
        "Targets: South Africa (1996) Growth, Employment and Redistribution: A "
        "Macroeconomic Strategy, Department of Finance, Integrated Scenario "
        "Projections table. The CPI inflation target for 1996\u201398 has been "
        "corrected here from the values given in Nattrass (2001) Table 1 "
        "(8.4 / 10.9 / 9.6), which reproduce GEAR's discarded Base Scenario rather "
        "than the Integrated Scenario actually adopted; the verified Integrated "
        "Scenario figures (8.0 / 9.7 / 8.1) are used instead, with 7.6 supplied for "
        "2000 (blank in the original table).",
        "Actuals, 1996\u201399: Nattrass, N. (2001) 'High Productivity Now', Table 1, "
        "citing SARB Quarterly Bulletin, March 2000.",
        "Actuals, 2000: South African Reserve Bank (2001) Annual Economic Report "
        "2001. Non-agric. formal employment growth and the real bank rate are shown "
        "as \u2014 for 2000 because the report gives component figures (private- and "
        "public-sector employment separately; discrete repo-rate change dates) but "
        "no single stated full-year figure comparable to the other years; no value "
        "is estimated in their place. The 2000 deficit figure (2.0% of GDP) is a "
        "fiscal-year (2000/01), not calendar-year, figure and is not strictly "
        "comparable to 1996\u201399.",
        "\u2014 denotes a figure not available from the cited source for that year "
        "(matches Nattrass's own convention in Table 1).",
    ]
    for note in notes:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(note)
        run.font.size = Pt(9)

    doc.save(OUT_DOCX)
    print(f"Saved docx table to: {OUT_DOCX}")


if __name__ == "__main__":
    main()
